import pandas as pd
import numpy as np
from sklearn.linear_model import LogisticRegression
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Alignment


# ============================================================
# 1. LOAD DATASET
# ============================================================

df = pd.read_csv("BLC-10.csv")

print(f"Dataset successfully loaded. Total rows: {len(df)}")


# ============================================================
# 2. PREPROCESS FEATURES
# ============================================================

df["Reviews_Cleaned"] = pd.to_numeric(
    df["Reviews"],
    errors="coerce"
).fillna(0.0)

df["Comments_Cleaned"] = pd.to_numeric(
    df["Comments"],
    errors="coerce"
).fillna(0.0)


# ============================================================
# 3. CLEAN AND NORMALIZE BUSINESS CATEGORIES
# ============================================================

df["Category"] = (
    df["Category"]
    .fillna("Unknown")
    .str.strip()
    .str.title()
)


# ============================================================
# 4. DEFINE TARGET LABEL
#    ORIGINAL LOGIC
# ============================================================

median_comments = df["Comments_Cleaned"].median()

df["Target_Label"] = (
    (df["Reviews_Cleaned"] >= 4.5) &
    (df["Comments_Cleaned"] >= median_comments)
).astype(int)


# ============================================================
# 5. SELECT FEATURES
#    REVIEWS + COMMENTS
# ============================================================

X = df[
    [
        "Reviews_Cleaned",
        "Comments_Cleaned"
    ]
]

y = df["Target_Label"]


# ============================================================
# 6. TRAIN LOGISTIC REGRESSION MODEL
# ============================================================

model = LogisticRegression(
    random_state=42
)

model.fit(X, y)

print(
    "Predictive model successfully trained "
    "on review and comment signals."
)


# ============================================================
# 7. PREDICT PROBABILITY
# ============================================================

df["Predicted_Probability"] = (
    model.predict_proba(X)[:, 1]
)


# ============================================================
# 8. SORT BY PREDICTED PROBABILITY
# ============================================================

df_sorted = df.sort_values(
    by=[
        "Predicted_Probability",
        "Reviews_Cleaned",
        "Comments_Cleaned"
    ],
    ascending=[
        False,
        False,
        False
    ]
).reset_index(drop=True)


# ============================================================
# 9. SPLIT INDIAN AND FOREIGN BUSINESSES
# ============================================================

indian_businesses = df_sorted[
    df_sorted["Country"]
    .astype(str)
    .str.strip()
    .str.lower()
    == "india"
].copy()


foreign_businesses = df_sorted[
    df_sorted["Country"]
    .astype(str)
    .str.strip()
    .str.lower()
    != "india"
].copy()


# ============================================================
# 10. FINAL OUTPUT COLUMNS
# ============================================================

columns_to_export = [
    "S.No",
    "Country",
    "City",
    "Company",
    "CONTACT",
    "Google Map",
    "Category",
    "Predicted_Probability"
]


indian_final = indian_businesses[
    columns_to_export
].copy()

foreign_final = foreign_businesses[
    columns_to_export
].copy()


# ============================================================
# 11. KEEP CONTACT AS TEXT
# ============================================================

indian_final["CONTACT"] = (
    indian_final["CONTACT"].astype(str)
)

foreign_final["CONTACT"] = (
    foreign_final["CONTACT"].astype(str)
)


# ============================================================
# 12. EXPORT TO EXCEL
# ============================================================

output_filename = "Country_Leads.xlsx"

with pd.ExcelWriter(
    output_filename,
    engine="openpyxl"
) as writer:

    indian_final.to_excel(
        writer,
        sheet_name="Indian Businesses",
        index=False
    )

    foreign_final.to_excel(
        writer,
        sheet_name="Foreign Businesses",
        index=False
    )


# ============================================================
# 13. FORMAT EXCEL FILE
# ============================================================

wb = load_workbook(output_filename)

for ws in wb.worksheets:

    probability_column = None
    contact_column = None

    # Find required columns
    for cell in ws[1]:

        if cell.value == "Predicted_Probability":
            probability_column = cell.column

        if cell.value == "CONTACT":
            contact_column = cell.column


    # --------------------------------------------------------
    # Format probability as percentage
    # --------------------------------------------------------

    if probability_column is not None:

        for row in range(2, ws.max_row + 1):

            cell = ws.cell(
                row=row,
                column=probability_column
            )

            cell.number_format = "0.00%"


    # --------------------------------------------------------
    # Format CONTACT as text
    # --------------------------------------------------------

    if contact_column is not None:

        for row in range(2, ws.max_row + 1):

            cell = ws.cell(
                row=row,
                column=contact_column
            )

            cell.number_format = "@"


    # --------------------------------------------------------
    # Header formatting
    # --------------------------------------------------------

    for cell in ws[1]:

        cell.font = Font(
            bold=True,
            color="FFFFFF"
        )

        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="595959"
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )


    # --------------------------------------------------------
    # Auto-adjust column width
    # --------------------------------------------------------

    for column_cells in ws.columns:

        max_length = 0

        column_letter = (
            column_cells[0].column_letter
        )

        for cell in column_cells:

            try:

                cell_length = len(
                    str(cell.value)
                )

                if cell_length > max_length:
                    max_length = cell_length

            except:
                pass

        ws.column_dimensions[
            column_letter
        ].width = min(
            max_length + 2,
            50
        )


# ============================================================
# 14. SAVE FINAL FILE
# ============================================================

wb.save(output_filename)


# ============================================================
# 15. FINAL OUTPUT
# ============================================================

print("\n==========================================")
print("Excel File Successfully Created")
print("==========================================")

print(f"Output file: {output_filename}")

print(
    f"\nIndian Businesses  : "
    f"{len(indian_final)}"
)

print(
    f"Foreign Businesses : "
    f"{len(foreign_final)}"
)

print(
    f"Total Businesses   : "
    f"{len(indian_final) + len(foreign_final)}"
)

print("\nSheets created:")
print("1. Indian Businesses")
print("2. Foreign Businesses")

print(
    "\nPredicted_Probability is displayed "
    "as a percentage."
)